"""Fill the user-supplied AI Builder Cup template; preserve mandatory section order."""
import argparse,hashlib,json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
ROOT=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser(description=__doc__);p.add_argument('--template',type=Path,required=True);p.add_argument('--out',type=Path,default=Path('/tmp/lingolive-submission-current/deck/LingoLive-Pitch-Deck.pptx'));p.add_argument('--public-repo',action='store_true');a=p.parse_args();r=Presentation(a.template)
assert len(r.slides)==16
# Remove only the template download-instructions page and the empty spare page.
for i in [14,1]:
 sid=r.slides._sldIdLst[i];r.part.drop_rel(sid.rId);r.slides._sldIdLst.remove(sid)
slides=list(r.slides);created=[]
BLUE='1A73E8';GREEN='188038';INK='202124';GRAY='5F6368';LIGHT='F1F6FD'
BASE='https://storage.googleapis.com/lingolive-demo-358650584819/ai-builder-cup-2026/submission-2026-10-18/'
def text(s,x,y,w,h,content,size=15,color=INK,bold=False,fill=None,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.06);tf.margin_right=Inches(.06);tf.margin_top=Inches(.03);tf.margin_bottom=0
 for i,line in enumerate(content.split('\n')):
  para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.text=line;para.font.name='Arial';para.font.size=Pt(size);para.font.bold=bold;para.font.color.rgb=RGBColor.from_string(color);para.space_after=Pt(7)
  if align is not None:para.alignment=align
 if fill:sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background()
 created.append((s,sh));return sh

def card(s,x,y,w,h,title,body):
 box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));box.fill.solid();box.fill.fore_color.rgb=RGBColor.from_string(LIGHT);box.line.color.rgb=RGBColor.from_string('DADCE0')
 if h<1:
  text(s,x+.13,y+.03,w-.26,.25,title,13,BLUE,True);text(s,x+.13,y+.29,w-.26,h-.31,body,11)
 else:
  text(s,x+.13,y+.10,w-.26,.36,title,15,BLUE,True);text(s,x+.13,y+.53,w-.26,h-.57,body,13)

def title(s,value):
 sh=next(sh for sh in s.shapes if sh.has_text_frame and sh.text.strip());sh.text=value;sh.height=Inches(.62)
 for q in sh.text_frame.paragraphs:q.font.name='Arial';q.font.size=Pt(18);q.font.bold=True;q.font.color.rgb=RGBColor.from_string(INK)

def link(s,y,label,url):
 sh=text(s,.55,y,8.8,.50,label,15,BLUE);sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url

cover=slides[0].shapes[0];cover.text='Team Details\nProject: LingoLive — Understand. Reply. Keep going.\nTeam name: CreativeEngine\nTeam leader: Ashish T. Vasant | Member: Janhvi Singh\nProblem: Help immigrants and interstate newcomers respond in an unfamiliar language.\nTheme: Media, Content & Digital Experiences | Category: Education'
for i,q in enumerate(cover.text_frame.paragraphs):q.font.name='Arial';q.font.size=Pt(18 if i==0 else 13);q.font.bold=i==0;q.space_after=Pt(5)

s=slides[1];title(s,'Brief about the idea');text(s,.5,1.52,8.9,.7,'Find your words in a new city.',30,INK,True);text(s,.5,2.36,8.9,.84,'A newcomer in Bengaluru needs to answer a cab driver in Kannada. LingoLive helps them understand, choose a reply and say it aloud.',17)
for x,t,b in [(0.55,'Understand','Live multilingual speech → meaning'),(3.65,'Reply','Useful options + pronunciation'),(6.75,'Learn later','Saved words and personal practice')]:card(s,x,3.55,2.75,1.25,t,b)

s=slides[2];title(s,'Opportunities');card(s,.5,1.55,9,1.00,'How is it different?','Connects the next turn, a pre-task plan, relevant personal/place memory, and later learning in one small interface.');card(s,.5,2.72,9,1.00,'How does it solve the problem?','Prioritizes usable replies during the encounter; the learner reviews word meanings afterwards. Repeated replies do not become new questions.');card(s,.5,3.89,9,1.05,'USP and intended users','Local, editable context follows the learner. Initial pilots: relocation teams, universities and community language groups. Outcomes and demand are not yet measured.')

s=slides[3];title(s,'List of features offered by the solution')
features=[('Hear and respond','Google multilingual transcription; three short replies.'),('Say it your way','Express familiar-language intent in the learning language.'),('Recognize readbacks','Fuzzy repeat guards + session voice estimates; correction.'),('Plan before leaving','Cab → optional tea → station; explicit progress controls.'),('Relevant memory','Editable routines; opt-in saved-place notes on return.'),('Learn afterwards','Saved review, word bubbles and pronunciation.'),('Minimal interface','One large microphone control; typing and first-use tutorial.'),('Bounded access','Private host code or your own Google key; usage guards.')]
for i,(t,b) in enumerate(features):card(s,.5+(i%2)*4.6,1.5+(i//2)*.90,4.4,.79,t,b)

s=slides[4];title(s,'Process flow diagram or Use-case diagram')
for x,t,b in [(0.5,'1. Listen','16 kHz browser microphone'),(3.7,'2. Transcribe','Gemini Transcribe Live'),(6.9,'3. Route the turn','Text/voice repeat estimates')]:card(s,x,1.55,2.6,1.25,t,b)
text(s,3.15,1.88,.45,.5,'→',23,BLUE,True);text(s,6.35,1.88,.45,.5,'→',23,BLUE,True)
card(s,.5,3.2,4.35,1.35,'Your repeated words','Keep in local conversation context. Wait for the other speaker. Uncertain meaning → confirm the speaker.');card(s,5.15,3.2,4.35,1.35,'Other / unknown turn','Two Flash-Lite branches → meaning + replies → learner speaks → Saved learning later.');text(s,.55,4.82,8.9,.35,'Speaker estimates help routing; they are not reliable identity. Explicit correction always wins.',11,GRAY)

s=slides[5];title(s,'Wireframes/Mock diagrams of the proposed solution (optional)')
for x,t,b in [(0.5,'Conversation','Heard words\nMeaning in your language\nThree usable replies\nOriginal / Pronunciation\nLarge listen button'),(3.65,'Before the task','Where do you want to go?\nOptional stops\nLearner-controlled progress\nRelevant local memories'),(6.8,'Saved learning','Your conversation library\nTap a word → meaning\nPronunciation and practice\nNo new lesson request')]:card(s,x,1.65,2.7,2.8,t,b)
text(s,.55,4.8,8.8,.3,'Functional prototype screens are shown in the snapshots section and video.',12,GRAY)

s=slides[6];title(s,'Architecture diagram of the proposed solution')
card(s,.5,1.6,2.65,2.8,'Browser','AudioWorklet capture\nLocal repeat / voice guard\nConversation + plan context\nIndexedDB saved lessons\nLocal, opt-in place memory')
card(s,3.65,1.6,2.65,2.8,'Google Cloud Run','FastAPI HTTP + WebSocket\nHost code / personal-key session\nRequest + audio quotas\nExact-source validation\nSecret Manager access code')
card(s,6.8,1.6,2.65,2.8,'Google AI','Gemini 3.5 Transcribe Live\nGemini 3.1 Flash-Lite\nSource and reply branches\nVertex service account\nBYOK: Gemini Developer API')
text(s,3.18,2.5,.4,.5,'⇄',22,BLUE,True);text(s,6.32,2.5,.4,.5,'⇄',22,BLUE,True);text(s,.55,4.7,8.9,.5,'Firebase / Firestore are not used in this prototype. Any future cloud sync would require a separate opt-in design.',12,GRAY)

s=slides[7];title(s,'Technologies to be used in the solution')
for i,(t,b) in enumerate([('Google AI','Gemini 3.5 Transcribe Live; two parallel Gemini 3.1 Flash-Lite branches with minimal thinking.'),('Google Cloud','Cloud Run, Vertex AI, Secret Manager, Cloud Build and Artifact Registry.'),('Web and backend','Responsive JavaScript PWA, AudioWorklet, IndexedDB, Python FastAPI and WebSockets.'),('Local adaptation','Bounded context and lexical/acoustic heuristics; no model-weight fine-tuning or stored voiceprint.')]):card(s,.5,1.55+i*.86,9,.76,t,b)

s=slides[8];title(s,'Estimated implementation cost (optional)');text(s,.55,1.55,8.8,.5,'Cost model: pay for useful turns, bound prototype usage.',21,INK,True)
text(s,.55,2.23,8.8,2.6,'• A 15-minute listening session uses up to 15 minutes of streaming audio.\n• N new lessons use 2 × N Flash-Lite generation calls.\n• Detected repeats, saved word review and script switching add no lesson call.\n• Cloud Run: 1 CPU / 512 MiB, request billing, min 0 / max 1 instance.\n• Host limits: 400 raw model calls/day and 60 audio minutes/day.\n• Prices depend on tokens, provider rates and actual runtime; no dollar estimate is asserted.',15)
text(s,.55,4.95,8.8,.32,'Quotas are prototype safeguards, not a hard billing cap. A wider launch needs durable per-user budgets.',11,GRAY)

s=slides[9];title(s,'Snapshots of the prototype')
shots=ROOT/'validation';names=[('web-plans-cloud.png','Plan before leaving'),('web-express-cloud.png','Find a useful reply'),('web-replies-review.png','Learn afterwards')]
from PIL import Image
for i,(name,label) in enumerate(names):
 file=shots/name;im=Image.open(file);im=im.crop((0,0,im.width,min(im.height,round(im.width*1.95))))
 tmp=a.out.parent/(name+'.crop.png');tmp.parent.mkdir(parents=True,exist_ok=True);im.save(tmp)
 s.shapes.add_picture(str(tmp),Inches(1.15+i*3.15),Inches(1.55),height=Inches(3.1));text(s,.5+i*3.15,4.66,2.8,.34,label,13,BLUE,True,align=PP_ALIGN.CENTER)
text(s,.55,5.05,8.9,.25,'Recorded prototype views. Word-review screenshot uses labelled test data; the video includes real Google responses.',10,GRAY)

s=slides[10];title(s,'Prototype Performance report/Benchmarking');text(s,.55,1.55,8.8,.5,'Functional evidence, with clear measurement limits.',22,INK,True)
rows=[('Earlier v1.5 warm HTTP lessons','10 complete lessons · mean 1.79 s · historical sample'),('Later v1.9 typed HTTP checks','4 functional cases · mean 2.065 s · not a latency guarantee'),('Current readback guard','25 local checks + 10 browser-flow checks; repeats add no lesson'),('Voice estimates','3/4 held-out clean synthetic learner clips recognized; 0/4 others labelled user'),('Service checks','112 backend tests; real Google replies and current Cloud Run guards verified')]
for i,(t,b) in enumerate(rows):text(s,.55,2.2+i*.51,3.0,.46,t,13,BLUE,True);text(s,3.62,2.2+i*.51,5.85,.46,b,12)
text(s,.55,4.95,8.9,.3,'No measured learning-outcome, noisy-phone, cold-start or population speaker-accuracy claim.',11,GRAY)

s=slides[11];title(s,'Additional Details/Future Development (if any)');card(s,.5,1.6,4.35,2.95,'Social value and business case','Help newcomers participate in everyday encounters.\nPilot with relocation teams, universities and community language groups.\nProposed organization-sponsored access; validate demand before expansion.');card(s,5.15,1.6,4.35,2.95,'Next engineering work','Representative mobile/noisy-language studies.\nDurable per-user budgets and stronger identity.\nOpt-in cross-device memory, potentially Firestore.\nMeasure task completion, reply usefulness and learning retention.');text(s,.55,4.88,8.9,.35,'Current limits: ASR can mishear; voice estimates can be wrong; local profiles are not private cloud accounts.',11,GRAY)

s=slides[12];title(s,'Provide links to your:')
link(s,1.6,'GitHub source — github.com/U4AR/LingoLive','https://github.com/U4AR/LingoLive')
link(s,2.35,'Demo video — 3–4 minute working demonstration',BASE+'LingoLive-Demo.mp4')
link(s,3.1,'Final product — live Google Cloud Run prototype','https://lingolive-358650584819.us-central1.run.app')
link(s,3.85,'Submission materials — pitch, technical PDF and entry text',BASE+'index.html')
text(s,.55,4.6,8.9,.6,'Keep the source repository publicly accessible throughout evaluation.'+'\nAI access: use your Google key, or obtain the host code privately from the team. No keys or codes are included.',11,GRAY)

# Closing branding is unmodified. Add restrained pagination to content slides.
for i,s in enumerate(slides[1:-1],2):text(s,8.95,5.26,.5,.2,str(i),8,GRAY,align=PP_ALIGN.RIGHT)
for s,sh in created:assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=r.slide_width+5 and sh.top+sh.height<=r.slide_height+5
assert len(slides)==14;a.out.parent.mkdir(parents=True,exist_ok=True);r.save(a.out)
report={'source_template_sha256':hashlib.sha256(a.template.read_bytes()).hexdigest(),'source_slides':16,'output_slides':14,'removed_pages':['download instruction page','empty spare page'],'required_sections_preserved_in_order':True,'branding_preserved':True,'team':'CreativeEngine','team_members':['Ashish T. Vasant','Janhvi Singh'],'repo_public_at_render':a.public_repo,'file':a.out.name,'bytes':a.out.stat().st_size}
(a.out.parent/'pitch-template-verification.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report))
